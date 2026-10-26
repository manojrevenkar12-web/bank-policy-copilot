"""
Reads policy documents from Word, PDF and wiki HTML exports into one structure: header fields,
sections, numbered paragraphs (Tz.). A document without the header fields that make it citable
(number, version, date of effect, classification) is rejected, not guessed.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader

LEVELS = {"öffentlich": 1, "intern": 2, "vertraulich": 3, "streng vertraulich": 4}
REQUIRED = ("Dokumentnummer", "Version", "Gültig ab", "Klassifizierung")

_TZ = re.compile(r"^Tz\.\s*(\d+(?:\.\d+)*)\s+(.*)$", re.S)
_HEADING = re.compile(r"^(\d+)\s+([A-ZÄÖÜ].{2,80})$")
_FIELD = re.compile(r"^(Dokumentnummer|Version|Gültig ab|Klassifizierung|Verantwortlich|Bezug):\s*(.+)$")


class ParseError(ValueError):
    pass


@dataclass
class ParsedParagraph:
    tz: str
    section: str
    text: str


@dataclass
class ParsedDocument:
    title: str
    fields: dict[str, str]
    paragraphs: list[ParsedParagraph] = field(default_factory=list)

    @property
    def doc_key(self) -> str:
        return self.fields["Dokumentnummer"]

    @property
    def version(self) -> int:
        return int(self.fields["Version"])

    @property
    def valid_from(self) -> date:
        d, m, y = (int(x) for x in self.fields["Gültig ab"].split("."))
        return date(y, m, d)

    @property
    def classification(self) -> int:
        level = LEVELS.get(self.fields["Klassifizierung"].strip().lower())
        if level is None:
            raise ParseError(f"unknown classification {self.fields['Klassifizierung']!r}")
        return level

    @property
    def references(self) -> list[str]:
        return [r.strip() for r in self.fields.get("Bezug", "").split(",") if r.strip()]

    def validate(self) -> None:
        missing = [f for f in REQUIRED if not self.fields.get(f)]
        if missing:
            raise ParseError(f"missing header fields: {', '.join(missing)}")
        if not self.paragraphs:
            raise ParseError("no numbered paragraphs (Tz.) found")
        _ = self.valid_from, self.classification, self.version
        seen: set[str] = set()
        for p in self.paragraphs:
            if p.tz in seen:
                raise ParseError(f"paragraph Tz. {p.tz} appears twice")
            seen.add(p.tz)


def _lines_to_paragraphs(lines: list[str]) -> list[ParsedParagraph]:
    """Groups lines into Tz. paragraphs; a line that is neither a heading nor a new Tz. continues
    the previous paragraph (PDF text wraps)."""
    out: list[ParsedParagraph] = []
    section = ""
    for raw in lines:
        line = " ".join(raw.split())
        if not line:
            continue
        m = _TZ.match(line)
        if m:
            out.append(ParsedParagraph(m.group(1), section, m.group(2).strip()))
            continue
        h = _HEADING.match(line)
        if h and not line.endswith("."):
            section = h.group(2).strip()
            continue
        if out:
            out[-1].text = f"{out[-1].text} {line}"
    return out


def parse_docx(path: Path) -> ParsedDocument:
    d = DocxDocument(str(path))
    fields: dict[str, str] = {}
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2 and cells[0]:
                fields[cells[0]] = cells[1]
    title = ""
    lines: list[str] = []
    for para in d.paragraphs:
        style = para.style.name if para.style is not None else ""
        text = para.text.strip()
        if style == "Title" and not title:
            title = text
        elif style.startswith("Heading"):
            lines.append(text)
        else:
            lines.append(text)
    return ParsedDocument(title, fields, _lines_to_paragraphs(lines))


def parse_pdf(path: Path) -> ParsedDocument:
    reader = PdfReader(str(path))
    raw = "\n".join(page.extract_text() or "" for page in reader.pages)
    lines = [line.strip() for line in raw.splitlines()]
    fields: dict[str, str] = {}
    title_parts: list[str] = []
    body: list[str] = []
    in_header = True
    for line in lines:
        m = _FIELD.match(line)
        if m:
            fields[m.group(1)] = m.group(2).strip()
            in_header = False
            continue
        if in_header and not fields:
            title_parts.append(line)
        else:
            body.append(line)
    return ParsedDocument(" ".join(t for t in title_parts if t), fields, _lines_to_paragraphs(body))


def parse_html(path: Path) -> ParsedDocument:
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
    title = soup.find("h1") or soup.find("title")
    fields: dict[str, str] = {}
    for row in soup.select("table.meta tr"):
        th, td = row.find("th"), row.find("td")
        if th and td:
            fields[th.get_text(strip=True)] = td.get_text(strip=True)
    lines: list[str] = []
    for el in soup.find_all(["h2", "p"]):
        lines.append(el.get_text(" ", strip=True))
    return ParsedDocument(title.get_text(strip=True) if title else "", fields, _lines_to_paragraphs(lines))


def parse(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        doc = parse_docx(path)
    elif suffix == ".pdf":
        doc = parse_pdf(path)
    elif suffix in (".html", ".htm"):
        doc = parse_html(path)
    else:
        raise ParseError(f"unsupported file type {suffix}")
    doc.validate()
    return doc
